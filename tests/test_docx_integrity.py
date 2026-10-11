"""Observable DOCX preservation and diagram-fidelity regressions."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
TOOLS = ROOT / 'tools/docx'
sys.path.insert(0, str(TOOLS))
import build_docx as builder
import render_docx as renderer
import compare_docx as comparison
import verify_sources as sources
import draw_sequence as sequence

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


class IntegrityTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix='icode-docx-integrity-')
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)
        self.source = self.base / 'guide.md'
        self.docx = self.base / 'guide.docx'
        self.manifest = self.base / 'guide.manifest.json'
        Image.new('RGB', (40, 40), 'red').save(self.base / 'figure.png')

    def build(self, text):
        self.source.write_text(text, encoding='utf-8')
        return builder.build(self.source, self.docx, self.manifest)

    def check(self, *extra):
        result = subprocess.run([sys.executable, str(TOOLS / 'inspect_docx.py'), str(self.docx),
                                 '--source', str(self.source), '--manifest', str(self.manifest),
                                 *extra], capture_output=True, text=True)
        return result, json.loads(result.stdout) if result.stdout.strip() else None

    def update_output_hash(self):
        data = json.loads(self.manifest.read_text())
        data['output']['sha256'] = builder.sha256_file(self.docx)
        self.manifest.write_text(json.dumps(data))

    def test_valid_build_and_repeated_image_references_pass(self):
        self.build('# 原理\n\n## SDK流程\n\n![图1](figure.png)\n\n![图2](figure.png)\n')
        result, report = self.check()
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        self.assertTrue(report['passed'])

    def test_tall_image_fits_printable_height_and_keeps_aspect_ratio(self):
        Image.new('RGB', (200, 1600), 'red').save(self.base / 'figure.png')
        self.build('# 原理\n\n![长图](figure.png)\n')
        document = Document(self.docx)
        shape = document.inline_shapes[0]
        section = document.sections[0]
        self.assertLessEqual(shape.height, section.page_height - section.top_margin - section.bottom_margin)
        self.assertAlmostEqual(shape.width / shape.height, 200 / 1600, places=5)

    def test_changed_source_fails_even_if_structure_is_identical(self):
        self.build('# 原理\n\n相机的时间。\n')
        self.source.write_text('# 原理\n\n电脑的时间。\n')
        result, _ = self.check()
        self.assertEqual(result.returncode, 1)

    def test_changed_word_fails_output_hash_check(self):
        self.build('# 原理\n\n原始正文。\n')
        document = Document(self.docx)
        document.paragraphs[1].text = '替换正文。'
        document.save(self.docx)
        result, _ = self.check()
        self.assertEqual(result.returncode, 1)

    def test_removed_heading_is_detected_even_with_valid_output_hash(self):
        self.build('# 原理\n\n## 时间转换\n\n时间单位是微秒。\n')
        document = Document(self.docx)
        document.paragraphs[1].style = document.styles['Normal']
        document.save(self.docx)
        self.update_output_hash()
        result, _ = self.check()
        self.assertEqual(result.returncode, 1)

    def test_changed_source_image_is_detected(self):
        self.build('# 原理\n\n![图1](figure.png)\n')
        Image.new('RGB', (40, 40), 'blue').save(self.base / 'figure.png')
        result, _ = self.check()
        self.assertEqual(result.returncode, 1)

    def test_broken_drawing_relationship_is_detected(self):
        self.build('# 原理\n\n![图1](figure.png)\n')
        with zipfile.ZipFile(self.docx) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        xml = ET.fromstring(parts['word/document.xml'])
        blip = xml.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
        blip.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', 'missing')
        parts['word/document.xml'] = ET.tostring(xml)
        with zipfile.ZipFile(self.docx, 'w') as archive:
            for name, payload in parts.items():
                archive.writestr(name, payload)
        self.update_output_hash()
        result, _ = self.check()
        self.assertEqual(result.returncode, 1)

    def test_markdown_syntax_inside_code_is_not_counted_as_document_structure(self):
        self.build('# 原理\n\n```text\n## 不是标题\n![不是图](missing.png)\n| a | b |\n|---|---|\n```\n')
        result, report = self.check()
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        self.assertEqual(report['expected']['images'], 0)

    def test_unsupported_mermaid_is_preserved_and_warned(self):
        self.build('# 时序\n\n```mermaid\nsequenceDiagram\nparticipant A\nparticipant B\nA->>B: 设置参数\n```\n')
        result, report = self.check()
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        self.assertTrue(report['warnings'])
        self.assertIn('设置参数', '\n'.join(p.text for p in Document(self.docx).paragraphs))

    def test_repeated_fallbacks_cannot_hide_a_removed_code_diagram(self):
        block = '```mermaid\nsequenceDiagram\nA->>B: 设置参数\n```\n'
        self.build('# 时序\n\n' + block + '\n' + block)
        document = Document(self.docx)
        paragraph = document.paragraphs[-1]._element
        paragraph.getparent().remove(paragraph)
        document.save(self.docx)
        self.update_output_hash()
        result, _ = self.check()
        self.assertEqual(result.returncode, 1)

    def test_require_diagrams_rejects_a_code_fallback(self):
        self.build('# 时序\n\n```mermaid\nsequenceDiagram\nA->>B: 设置参数\n```\n')
        result, _ = self.check('--require-diagrams')
        self.assertEqual(result.returncode, 1)

    def test_malformed_manifest_is_a_controlled_failure(self):
        self.build('# 原理\n')
        for payload in ('{', '[]', '{"input": null}'):
            with self.subTest(payload=payload):
                self.manifest.write_text(payload)
                result, _ = self.check()
                self.assertEqual(result.returncode, 1)
                self.assertNotIn('Traceback', result.stderr)

    def test_malformed_render_manifest_is_a_controlled_failure(self):
        self.build('# 原理\n')
        self.manifest.write_text('[]')
        result = subprocess.run([sys.executable, str(TOOLS / 'render_docx.py'), str(self.docx), str(self.base / 'preview'),
                                 '--manifest', str(TOOLS / 'renderer_manifest.json'), '--result-manifest', str(self.manifest)],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 1)
        self.assertNotIn('Traceback', result.stderr)

    def test_mermaid_does_not_silently_drop_conditions_shapes_or_styles(self):
        cases = ('flowchart LR\nA -->|成功| B',
                 'flowchart LR\nA{是否成功} --> B',
                 'flowchart LR\nA --> B --> C',
                 'flowchart LR\nA --> B\nclassDef sdk fill:#123456\nclass B sdk',
                 'sequenceDiagram\nA --> B',
                 'flowchart LR\nA -.-> B',
                 'flowchart LR\nA --> B\nB --> C\nC --> A',
                 'flowchart LR\nA --> B\nA --> C')
        for source in cases:
            with self.subTest(source=source):
                rendered, _ = builder.render_mermaid(source, self.base / 'diagram.png')
                self.assertFalse(rendered)

    def test_reverse_directions_place_first_node_on_requested_side(self):
        for direction in ('RL', 'BT'):
            with self.subTest(direction=direction):
                output = self.base / (direction + '.png')
                rendered, _ = builder.render_mermaid(f'flowchart {direction}\nA[A] --> B[B]', output)
                self.assertTrue(rendered)
                # Capture label positions rather than comparing a font-dependent bitmap.
                from unittest.mock import patch
                with patch('PIL.ImageDraw.ImageDraw.multiline_text') as draw:
                    builder.render_mermaid(f'flowchart {direction}\nA[A] --> B[B]', output)
                    points = {call.args[1]: call.args[0] for call in draw.call_args_list}
                axis = 0 if direction == 'RL' else 1
                self.assertGreater(points['A'][axis], points['B'][axis])

    def test_render_record_is_bound_to_word_and_page_hashes(self):
        self.build('# 原理\n')
        page = self.base / 'page-1.png'
        Image.new('RGB', (10, 10), 'white').save(page)
        renderer.update_manifest(self.manifest, 'passed', docx=self.docx, pages=[str(page)])
        qa = json.loads(self.manifest.read_text())['visual_qa']
        self.assertEqual(qa['docx_sha256'], builder.sha256_file(self.docx))
        self.assertEqual(qa['page_sha256'][str(page)], builder.sha256_file(page))

    def test_render_record_cannot_be_reused_after_word_changes(self):
        self.build('# 原理\n\n保留正文。\n')
        page = self.base / 'page-1.png'
        Image.new('RGB', (10, 10), 'white').save(page)
        renderer.update_manifest(self.manifest, 'passed', docx=self.docx, pages=[str(page)])
        document = Document(self.docx)
        document.paragraphs[1].text = '另一个正文。'
        document.save(self.docx)
        self.update_output_hash()
        result, report = self.check()
        self.assertEqual(result.returncode, 1)
        self.assertFalse(report['checks']['render_matches_word'])

    def test_render_record_cannot_be_reused_after_page_changes(self):
        self.build('# 原理\n')
        page = self.base / 'page-1.png'
        Image.new('RGB', (10, 10), 'white').save(page)
        renderer.update_manifest(self.manifest, 'passed', docx=self.docx, pages=[str(page)])
        Image.new('RGB', (10, 10), 'red').save(page)
        result, report = self.check()
        self.assertEqual(result.returncode, 1)
        self.assertFalse(report['checks']['render_pages_unchanged'])


class RevisionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix='icode-docx-revision-')
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)
        self.before = self.base / 'before.docx'
        self.after = self.base / 'after.docx'
        self.figure = self.base / 'figure.png'
        Image.new('RGB', (40, 40), 'red').save(self.figure)
        document = Document()
        document.add_heading('原理', level=1)
        document.add_paragraph('这段不能改变。')
        document.add_heading('5.1 空间坐标', level=2)
        document.add_paragraph('X向右，Y向下，Z向前，单位毫米。')
        document.add_picture(str(self.figure))
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '深度单位'
        table.cell(0, 1).text = '毫米'
        document.sections[0].header.paragraphs[0].text = '固定页眉'
        document.save(self.before)
        document.save(self.after)

    def test_local_rewrite_preserves_outside_and_lists_deleted_explanation(self):
        document = Document(self.after)
        document.paragraphs[3].text = '空间点的三个坐标都以毫米为单位。'
        document.save(self.after)
        report = comparison.compare(self.before, self.after, ['5.1 空间坐标'], [])
        self.assertTrue(report['passed'])
        self.assertIn('X向右，Y向下，Z向前，单位毫米。', report['removed_text'])

    def test_unrelated_body_edit_is_rejected(self):
        document = Document(self.after)
        document.paragraphs[1].text = '被误改的段落。'
        document.save(self.after)
        self.assertFalse(comparison.compare(self.before, self.after, ['5.1 空间坐标'], [])['passed'])

    def test_section_allowance_does_not_hide_image_replacement(self):
        document = Document(self.after)
        part = document.part.related_parts[document.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed]
        Image.new('RGB', (40, 40), 'blue').save(self.figure)
        part._blob = self.figure.read_bytes()
        document.save(self.after)
        self.assertFalse(comparison.compare(self.before, self.after, ['5.1 空间坐标'], [])['passed'])
        self.assertTrue(comparison.compare(self.before, self.after, ['5.1 空间坐标'], [1])['passed'])

    def test_unknown_and_ambiguous_section_names_are_rejected(self):
        with self.assertRaises(ValueError):
            comparison.compare(self.before, self.after, ['不存在'], [])
        for path in (self.before, self.after):
            document = Document(path)
            document.add_heading('5.1 空间坐标', level=2)
            document.save(path)
        with self.assertRaises(ValueError):
            comparison.compare(self.before, self.after, ['5.1 空间坐标'], [])

    def test_figure_count_change_and_header_edit_are_rejected(self):
        document = Document(self.after)
        document.add_picture(str(self.figure))
        document.sections[0].header.paragraphs[0].text = '误改页眉'
        document.save(self.after)
        report = comparison.compare(self.before, self.after, ['5.1 空间坐标'], [1])
        self.assertFalse(report['passed'])
        self.assertFalse(report['checks']['outside_figures_preserved'])
        self.assertFalse(report['checks']['headers_footers_preserved'])

    def test_table_change_is_reported(self):
        document = Document(self.after)
        document.tables[0].cell(0, 1).text = '微米'
        document.save(self.after)
        report = comparison.compare(self.before, self.after, [], [])
        self.assertFalse(report['passed'])
        self.assertIn('毫米', report['removed_text'])

    def test_scope_failure_and_report_only_have_distinct_exit_status(self):
        document = Document(self.after)
        document.paragraphs[1].text = '被改的段落。'
        document.save(self.after)
        command = [sys.executable, str(TOOLS / 'compare_docx.py'), str(self.before), str(self.after)]
        result = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(result.returncode, 1)
        result = subprocess.run(command + ['--report-only'], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0)
        self.assertFalse(json.loads(result.stdout)['passed'])


class SourceTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix='icode-docx-sources-')
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)
        self.code = self.base / 'device.cpp'
        self.code.write_text('// 单位是微秒\nvoid setTime() {}\n')
        self.record = self.base / 'sources.json'
        self.data = {'files': [{'path': 'device.cpp', 'sha256': builder.sha256_file(self.code),
                               'anchors': [{'start': 2, 'end': 2, 'contains': 'setTime'}]}]}
        self.write()

    def write(self):
        self.record.write_text(json.dumps(self.data))

    def test_valid_source_and_anchors_are_read_only(self):
        original = self.code.read_bytes(), self.record.read_bytes()
        self.assertTrue(sources.verify(self.record, self.base)['passed'])
        self.assertEqual(original, (self.code.read_bytes(), self.record.read_bytes()))

    def test_stale_hash_and_wrong_anchor_fail(self):
        self.code.write_text('// 单位是毫秒\nvoid setTime() {}\n')
        self.assertFalse(sources.verify(self.record, self.base)['passed'])
        self.data['files'][0]['sha256'] = builder.sha256_file(self.code)
        self.data['files'][0]['anchors'][0]['contains'] = 'missingSymbol'
        self.write()
        self.assertFalse(sources.verify(self.record, self.base)['passed'])

    def test_invalid_range_and_empty_records_fail(self):
        self.data['files'][0]['anchors'][0]['end'] = 99
        self.write()
        with self.assertRaises(ValueError):
            sources.verify(self.record, self.base)
        self.data = {'files': []}
        self.write()
        with self.assertRaises(ValueError):
            sources.verify(self.record, self.base)

    def test_path_escape_and_symlink_outside_root_fail(self):
        outside = self.base / 'outside'
        root = self.base / 'repo'
        root.mkdir()
        outside.write_text('void setTime() {}\n')
        (root / 'link.cpp').symlink_to(outside)
        for path in ('../outside', str(outside), 'link.cpp'):
            self.data['files'][0]['path'] = path
            self.write()
            with self.assertRaises(ValueError):
                sources.verify(self.record, root)

    def test_malformed_input_is_a_controlled_cli_failure(self):
        for value in ('{', '[]', '{"files":[null]}'):
            self.record.write_text(value)
            result = subprocess.run([sys.executable, str(TOOLS / 'verify_sources.py'), str(self.record),
                                     '--root', str(self.base)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 1)
            self.assertNotIn('Traceback', result.stderr)


class SequenceTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix='icode-docx-sequence-')
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)
        self.spec = self.base / 'sequence.json'
        self.palette = self.base / 'palette.json'
        self.output = self.base / 'sequence.png'
        self.data = {'title': '请求与返回', 'participants': [
            {'id': 'app', 'label': '应用程序', 'role': 'application'},
            {'id': 'sdk', 'label': 'SDK', 'role': 'sdk'}],
            'steps': [{'from': 'app', 'to': 'sdk', 'text': '设置参数'},
                      {'from': 'sdk', 'to': 'sdk', 'text': '检查参数'},
                      {'from': 'sdk', 'to': 'app', 'text': '返回结果', 'reply': True, 'condition': '检查成功'}]}
        self.colors = {'application': {'fill': '#FFE0D0', 'stroke': '#B54E16', 'text': '#47200D'},
                       'sdk': {'fill': '#D6E9FF', 'stroke': '#1463AD', 'text': '#0A345B'}}
        self.write()

    def write(self):
        self.spec.write_text(json.dumps(self.data, ensure_ascii=False))
        self.palette.write_text(json.dumps(self.colors))

    def test_sequence_preserves_steps_palette_and_self_calls(self):
        report = sequence.draw(self.spec, self.palette, self.output)
        self.assertEqual(report['steps'], self.data['steps'])
        self.assertEqual(report['palette_sha256'], builder.sha256_file(self.palette))
        with Image.open(self.output) as image:
            pixels = {color for _, color in image.getcolors(image.width * image.height)}
            self.assertIn((255, 224, 208), pixels)
            self.assertIn((214, 233, 255), pixels)
            self.assertGreater(image.height, 400)

    def test_unknown_modules_and_participants_are_not_invented(self):
        self.data['participants'][1]['role'] = 'new-module'
        self.write()
        with self.assertRaises(ValueError):
            sequence.draw(self.spec, self.palette, self.output)
        self.data['participants'][1]['role'] = 'sdk'
        self.data['steps'][0]['to'] = 'firmware'
        self.write()
        with self.assertRaises(ValueError):
            sequence.draw(self.spec, self.palette, self.output)

    def test_duplicate_participants_and_unsupported_blocks_are_rejected(self):
        self.data['participants'][1]['id'] = 'app'
        self.write()
        with self.assertRaises(ValueError):
            sequence.draw(self.spec, self.palette, self.output)
        self.data['participants'][1]['id'] = 'sdk'
        self.data['steps'][0]['parallel'] = True
        self.write()
        with self.assertRaises(ValueError):
            sequence.draw(self.spec, self.palette, self.output)

    def test_overlong_diagram_and_text_are_rejected_before_output(self):
        self.data['steps'] *= 4
        self.write()
        with self.assertRaises(ValueError):
            sequence.draw(self.spec, self.palette, self.output)
        self.assertFalse(self.output.exists())
        self.data['steps'] = [{'from': 'app', 'to': 'sdk', 'text': '过长说明' * 100}]
        self.write()
        with self.assertRaises(ValueError):
            sequence.draw(self.spec, self.palette, self.output)

    def test_sequence_can_be_embedded_and_pass_diagram_checks(self):
        sequence.draw(self.spec, self.palette, self.output)
        source = self.base / 'guide.md'
        source.write_text('# 一条命令\n\n![图1 收发时序](sequence.png)\n')
        output, manifest = self.base / 'guide.docx', self.base / 'guide.manifest.json'
        builder.build(source, output, manifest)
        result = subprocess.run([sys.executable, str(TOOLS / 'inspect_docx.py'), str(output), '--source', str(source),
                                 '--manifest', str(manifest), '--require-diagrams'], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr + result.stdout)


if __name__ == '__main__':
    unittest.main(verbosity=2)
