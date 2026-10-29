"""Real-package regression tests for strict PPT image editing."""
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "tools/ppt/scripts/build_pptx.py"
sys.path.insert(0, str(BUILD.parent))
sys.path.insert(0, str(ROOT / "tools/docx"))
from check_render_output import check_render_output


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class PptImagesTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name)
        self.old = self.base / "old.png"
        self.new = self.base / "new.png"
        Image.new("RGB", (80, 40), "red").save(self.old)
        Image.new("RGB", (40, 80), "blue").save(self.new)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
        prs.slides.add_slide(prs.slide_layouts[6])
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.pic = slide.shapes.add_picture(str(self.old), Inches(1), Inches(1), Inches(4), Inches(3))
        self.region = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(1), Inches(2), Inches(2))
        self.region.name = "DIAGRAM_SLOT"
        self.text = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(1))
        self.text.text = "preserve me"
        self.template = self.base / "template.pptx"
        prs.save(self.template)
        self.output = self.base / "result.pptx"
        self.edits = self.base / "edits.json"
        self.detail = self.base / "detail.json"
        self.detail.write_text(json.dumps({"pages": [{"slide_number": 2, "image_slots": [
            {"slot_id": "diagram", "address": {"shape_id": self.region.shape_id},
             "expected_placeholder": "DIAGRAM_SLOT"}]}]}), encoding="utf-8")

    def image(self, **extra):
        return {"slide": 2, "source": "new.png", "alt_text": "蓝色示意图", "fit": "contain",
                "address": {"shape_id": self.pic.shape_id}, "expected_sha256": sha(self.old), **extra}

    def run_build(self, images=None, extra=None, flags=(), env=None):
        data = {"selected_slides": [2, 1], "edits": [], "images": images or []}
        data.update(extra or {})
        self.edits.write_text(json.dumps(data), encoding="utf-8")
        return subprocess.run([sys.executable, "-B", str(BUILD), str(self.template), str(self.edits),
                               str(self.output), "--detail", str(self.detail), *flags],
                              text=True, capture_output=True, env=env)

    def previews(self):
        previews = self.base / "previews"
        previews.mkdir(exist_ok=True)
        for n in (1, 2):
            Image.new("RGB", (40, 30), "white").save(previews / f"slide-{n}.png")
        return previews

    def test_replacement_preserves_shape_id_z_order_text_and_manifest(self):
        result = self.run_build([self.image()])
        self.assertEqual(result.returncode, 0, result.stderr)
        slide = Presentation(self.output).slides[0]
        self.assertEqual([s.shape_id for s in slide.shapes], [self.pic.shape_id, self.region.shape_id, self.text.shape_id])
        pic = slide.shapes[0]
        self.assertEqual(hashlib.sha256(pic.image.blob).hexdigest(), sha(self.new))
        self.assertAlmostEqual(pic.width / pic.height, .5)
        self.assertEqual(slide.shapes[2].text, "preserve me")
        self.assertEqual(pic._element.nvPicPr.cNvPr.get("descr"), "蓝色示意图")
        manifest = json.loads(self.output.with_suffix(".manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["output"]["sha256"], sha(self.output))
        self.assertEqual(manifest["images"][0]["original_slide"], 2)
        self.assertEqual(manifest["images"][0]["output_slide"], 1)
        self.assertEqual(manifest["visual_review"], "pending")

    def test_detail_explicit_region_can_be_replaced(self):
        result = self.run_build([{"slide": 2, "source": "new.png", "slot_id": "diagram", "alt_text": "区域图"}])
        self.assertEqual(result.returncode, 0, result.stderr)
        slide = Presentation(self.output).slides[0]
        self.assertEqual(slide.shapes[1].shape_id, self.region.shape_id)
        self.assertEqual(hashlib.sha256(slide.shapes[1].image.blob).hexdigest(), sha(self.new))

    def test_explicit_address_can_target_a_declared_region(self):
        result = self.run_build([{"slide": 2, "source": "new.png", "address": {"shape_id": self.region.shape_id},
                                 "expected_placeholder": "DIAGRAM_SLOT", "alt_text": "区域图"}])
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_text_unknown_address_field_fails_dry_run(self):
        result = self.run_build(extra={"edits": [{"slide": 2, "address": {
            "shape_id": self.text.shape_id, "paragraph": 0, "typo": 1}, "new_text": "x"}]}, flags=("--dry-run",))
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.output.exists())

    def test_free_region_insertion_and_manifest_verification(self):
        result = self.run_build([{"slide": 2, "source": "new.png", "box_inches": [8, 4, 1, 2], "alt_text": "free"}])
        self.assertEqual(result.returncode, 0, result.stderr)
        previews = self.previews()
        manifest = self.output.with_suffix(".manifest.json")
        report = check_render_output(self.output, previews, manifest)
        self.assertEqual(report["status"], "passed", report)
        self.new.write_bytes(self.old.read_bytes())
        report = check_render_output(self.output, previews, manifest)
        self.assertEqual(report["status"], "failed")
        self.assertTrue(any("source" in x for x in report["errors"]))

    def test_manifest_checks_actual_geometry_alt_and_embedded_image(self):
        result = self.run_build([self.image()])
        self.assertEqual(result.returncode, 0, result.stderr)
        path = self.output.with_suffix(".manifest.json")
        original = json.loads(path.read_text(encoding="utf-8"))
        for field, value in [("geometry", {"left": 0}), ("alt_text", "wrong"), ("embedded_sha256", "0" * 64)]:
            manifest = json.loads(json.dumps(original))
            manifest["images"][0][field] = value
            path.write_text(json.dumps(manifest), encoding="utf-8")
            with self.subTest(field=field):
                self.assertEqual(check_render_output(self.output, self.previews(), path)["status"], "failed")

    def test_legacy_checker_call_still_accepts_valid_previews(self):
        result = self.run_build()
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(check_render_output(self.output, self.previews())["status"], "passed")

    def test_manifest_checker_reports_corrupt_package_without_exception(self):
        result = self.run_build([self.image()])
        self.assertEqual(result.returncode, 0, result.stderr)
        self.output.write_bytes(b"broken package")
        report = check_render_output(self.output, self.previews(), self.output.with_suffix(".manifest.json"))
        self.assertEqual(report["status"], "failed")
        self.assertTrue(report["errors"])

    def test_raster_svg_checker_detects_changed_fallback(self):
        if importlib.util.find_spec("diagram_svg") is None:
            self.skipTest("SVG shared component unavailable")
        source = self.base / "diagram.svg"
        source.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100"><rect width="200" height="100" fill="#2244aa"/></svg>', encoding="utf-8")
        result = self.run_build([self.image(source="diagram.svg")])
        self.assertEqual(result.returncode, 0, result.stderr)
        manifest_path = self.output.with_suffix(".manifest.json")
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        Path(manifest["diagrams"][0]["fallback_png"]).write_bytes(b"tampered fallback")
        report = check_render_output(self.output, self.previews(), manifest_path)
        self.assertEqual(report["status"], "failed")

    def test_svg_raster_and_dual_are_real_packages_with_pending_review(self):
        if importlib.util.find_spec("diagram_svg") is None or importlib.util.find_spec("diagram_ooxml") is None:
            self.skipTest("SVG shared components not installed in this checkout")
        source = self.base / "diagram.svg"
        source.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100"><rect width="200" height="100" fill="#2244aa"/></svg>', encoding="utf-8")
        dry_run = self.run_build([self.image(source="diagram.svg")], flags=("--dry-run",))
        self.assertEqual(dry_run.returncode, 0, dry_run.stderr)
        self.assertFalse((self.output.parent / (self.output.stem + ".assets")).exists())
        for mode in ("raster", "dual"):
            with self.subTest(mode=mode):
                result = self.run_build([self.image(source="diagram.svg")], flags=("--svg-mode", mode))
                self.assertEqual(result.returncode, 0, result.stderr)
                manifest_path = self.output.with_suffix(".manifest.json")
                manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
                item = manifest["diagrams"][0]
                fallback = Path(item["fallback_png"])
                self.assertEqual(fallback.parent, self.output.parent / (self.output.stem + ".assets") / "svg")
                self.assertTrue(fallback.is_file())
                self.assertEqual(item["source_format"], "svg")
                self.assertEqual(item["vector_preserved"], mode == "dual")
                self.assertEqual(item["validation"]["visual_review"], "pending")
                self.assertEqual(item["pixel_width"] / item["pixel_height"], 2)
                with zipfile.ZipFile(self.output) as package:
                    self.assertEqual(any(n.endswith(".svg") for n in package.namelist()), mode == "dual")
                report = check_render_output(self.output, self.previews(), manifest_path)
                self.assertEqual(report["status"], "passed", report)
                if mode == "dual":
                    self.assertEqual(item["consumer_compatibility"], "pending")
                    self.assertEqual(item.get("diagram_status"), "native_vector_pending")

    def test_svg_uses_configured_default_local_font(self):
        if importlib.util.find_spec("diagram_svg") is None:
            self.skipTest("SVG shared component unavailable")
        font = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/arial.ttf"
        if not font.is_file():
            self.skipTest("explicit local font fixture unavailable")
        source = self.base / "diagram.svg"
        source.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100"><text x="5" y="30" font-family="Arial">Diagram</text></svg>', encoding="utf-8")
        result = self.run_build([self.image(source="diagram.svg")], env={**os.environ, "ICODE_SVG_FONT_FILES": str(font)})
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_xml_invalid_alt_text_fails_dry_run(self):
        result = self.run_build([self.image(alt_text="bad\x01alt")], flags=("--dry-run",))
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.output.exists())

    def test_group_picture_replacement_preserves_group_and_siblings(self):
        prs = Presentation(self.template)
        slide = prs.slides[1]
        group = slide.shapes.add_group_shape()
        pic = group.shapes.add_picture(str(self.old), Inches(1), Inches(1), Inches(2), Inches(1))
        child = group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1))
        child.text = "group text"
        group_id, picture_id, child_id = group.shape_id, pic.shape_id, child.shape_id
        prs.save(self.template)
        result = self.run_build([self.image(address={"shape_id": picture_id})])
        self.assertEqual(result.returncode, 0, result.stderr)
        group = Presentation(self.output).slides[0].shapes[-1]
        self.assertEqual(group.shape_id, group_id)
        self.assertEqual([s.shape_id for s in group.shapes], [picture_id, child_id])
        self.assertEqual(group.shapes[1].text, "group text")
        self.assertEqual(hashlib.sha256(group.shapes[0].image.blob).hexdigest(), sha(self.new))

    def test_declared_group_cannot_discard_its_children(self):
        prs = Presentation(self.template)
        group = prs.slides[1].shapes.add_group_shape()
        group.name = "GROUP_SLOT"
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text = "keep child"
        prs.save(self.template)
        self.detail.write_text(json.dumps({"pages": [{"slide_number": 2, "image_slots": [{
            "slot_id": "group", "address": {"shape_id": group.shape_id}, "expected_placeholder": "GROUP_SLOT"}]}]}), encoding="utf-8")
        result = self.run_build([{"slide": 2, "slot_id": "group", "source": "new.png", "alt_text": "x"}])
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.output.exists())

    def test_dual_ambiguous_svg_sources_fail_dry_run(self):
        if importlib.util.find_spec("diagram_svg") is None or importlib.util.find_spec("diagram_ooxml") is None:
            self.skipTest("SVG shared components unavailable")
        svg = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100"><rect width="200" height="100" fill="#2244aa"/></svg>'
        (self.base / "one.svg").write_text(svg, encoding="utf-8")
        (self.base / "two.svg").write_text(svg.replace("<rect", "<!--same raster--><rect"), encoding="utf-8")
        # Identical box widths produce the identical fallback pixel dimensions.
        prs = Presentation(self.template)
        prs.slides[1].shapes[1].width = Inches(4)
        prs.save(self.template)
        result = self.run_build([self.image(source="one.svg"), {
            "slide": 2, "slot_id": "diagram", "source": "two.svg", "alt_text": "second"}],
            flags=("--svg-mode", "dual", "--dry-run"))
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.output.exists())

    def test_invalid_images_fail_before_output_even_dry_run(self):
        invalid = [
            self.image(typo=True), self.image(fit="cover"), self.image(alt_text=""),
            self.image(expected_sha256="0" * 64), self.image(slide=1),
            {"slide": 2, "source": "new.png", "alt_text": "x", "address": {"shape_id": self.text.shape_id}, "expected_placeholder": "preserve me"},
            {"slide": 2, "source": "new.png", "alt_text": "x", "box_inches": [0, 0, 11, 1]},
            {"slide": 2, "source": "new.png", "alt_text": "x", "box_inches": [1, 1, 1, 1]},
        ]
        for image in invalid:
            for flags in [(), ("--dry-run",)]:
                with self.subTest(image=image, flags=flags):
                    result = self.run_build([image], extra={"selected_slides": [2]}, flags=flags)
                    self.assertNotEqual(result.returncode, 0, result.stdout)
                    self.assertFalse(self.output.exists())
                    self.assertFalse(self.output.with_suffix(".manifest.json").exists())

    def test_duplicate_target_and_unselected_text_fail(self):
        self.assertNotEqual(self.run_build([self.image(), self.image()]).returncode, 0)
        result = self.run_build(extra={"selected_slides": [1], "edits": [{"slide": 2,
             "address": {"shape_id": self.text.shape_id}, "new_text": "x"}]}, flags=("--dry-run",))
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.output.exists())

    def test_output_cannot_overwrite_input(self):
        before = self.template.read_bytes()
        self.output = self.template
        self.assertNotEqual(self.run_build([self.image()]).returncode, 0)
        self.assertEqual(self.template.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
