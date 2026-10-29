"""Build-only qualification fixtures must preserve Chinese SVG and both modes."""
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from xml.etree import ElementTree as ET
import zipfile

from pptx import Presentation
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'tools/docx/qualify_svg.py'
SVG = '{http://www.w3.org/2000/svg}'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


class SvgQualificationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory(prefix='icode-svg-qualification-')
        cls.base = Path(cls.temporary.name)
        cls.output = cls.base / 'fixture'
        cls.result = None
        if SCRIPT.is_file():
            # Always exercise this checkout's entrypoint and sibling builders.
            # Inherited paths may supply installed runtime dependencies, but
            # are never required to find the repository's own tool modules.
            cls.entrypoint = SCRIPT
            cls.environment = os.environ.copy()
            paths = [str(ROOT / 'tools/docx'), str(ROOT / 'tools/ppt/scripts')]
            if cls.environment.get('PYTHONPATH'):
                paths.append(cls.environment['PYTHONPATH'])
            cls.environment['PYTHONPATH'] = os.pathsep.join(paths)
            candidates = [Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts/msyh.ttc',
                          Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')]
            cls.font = next((path for path in candidates if path.is_file()), candidates[0])
            cls.result = subprocess.run([sys.executable, '-X', 'utf8', str(cls.entrypoint),
                                         '--output-dir', str(cls.output), '--font', str(cls.font),
                                         '--font', str(cls.font)],
                                        capture_output=True, text=True, encoding='utf-8', env=cls.environment)

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def fixture(self):
        self.assertTrue(SCRIPT.is_file(), 'qualification entrypoint must exist')
        self.assertEqual(self.result.returncode, 0, self.result.stdout + self.result.stderr)
        return json.loads((self.output / 'qualification.json').read_text(encoding='utf-8'))

    def test_fixture_contains_chinese_sources_png_and_hashes_without_office_claim(self):
        report = self.fixture()
        self.assertEqual(report['status'], {'build': 'passed', 'libreoffice': 'pending', 'visual_review': 'pending'})
        self.assertIn('Microsoft Office/WPS optional', report['scope'])
        self.assertEqual(len(report['diagrams']), 3)
        descriptions = ' '.join(item['description'] for item in report['diagrams'])
        self.assertIn('IMU', descriptions)
        self.assertIn('时间', descriptions)
        self.assertIn('异常', descriptions)
        for item in report['diagrams']:
            source = Path(item['source'])
            self.assertTrue(source.is_relative_to(self.output))
            root = ET.parse(source).getroot()
            self.assertTrue(root.findall('.//' + SVG + 'text'))
            self.assertTrue(root.findall('.//' + SVG + 'marker'))
            text = ''.join(root.itertext())
            self.assertTrue(any('\u4e00' <= char <= '\u9fff' for char in text))
            self.assertFalse(root.findall('.//' + SVG + 'image'))
        records = report['artifacts']
        self.assertTrue(any(item['path'].endswith('.png') for item in records))
        self.assertTrue(any(item['path'].endswith('guide.md') for item in records))
        self.assertEqual(len(records), len(list(path for path in self.output.rglob('*') if path.is_file())) - 1)
        for item in records:
            path = Path(item['path'])
            self.assertTrue(path.is_relative_to(self.output))
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), item['sha256'])

    def test_four_packages_have_three_diagrams_and_only_dual_contains_svg(self):
        report = self.fixture()
        self.assertEqual(len(report['packages']), 4)
        for package in report['packages']:
            path, mode = Path(package['path']), package['mode']
            with self.subTest(package=path.name), zipfile.ZipFile(path) as archive:
                svg_parts = [name for name in archive.namelist() if name.endswith('.svg')]
                self.assertEqual(len(svg_parts), 3 if mode == 'dual' else 0)
                if path.suffix == '.docx':
                    root = ET.fromstring(archive.read('word/document.xml'))
                    self.assertEqual(len(root.findall('.//' + W + 'drawing')), 3)
                    self.assertEqual(len(root.findall('.//' + W + 'pageBreakBefore')), 2)
                else:
                    deck = Presentation(path)
                    self.assertEqual(len(deck.slides), 3)
                    for slide in deck.slides:
                        self.assertEqual(len(slide.shapes), 2)
                        self.assertEqual(slide.shapes[0].text_frame.paragraphs[0].runs[0].font.name, 'Noto Sans CJK SC')
                        image = slide.shapes[1]
                        self.assertGreaterEqual(image.left, 457200)
                        self.assertGreaterEqual(image.top, 1097280)
                manifest = json.loads(Path(package['manifest']).read_text(encoding='utf-8'))
                self.assertEqual(manifest['output']['sha256'], hashlib.sha256(path.read_bytes()).hexdigest())
                self.assertEqual(len(manifest['diagrams']), 3)
                self.assertTrue(all(Path(item['source']).is_relative_to(self.output) for item in manifest['diagrams']))

    def test_cli_refuses_existing_directory_without_modifying_it(self):
        self.fixture()
        marker = self.output / 'qualification.json'
        before = marker.read_bytes()
        result = subprocess.run([sys.executable, '-X', 'utf8', str(self.entrypoint),
                                 '--output-dir', str(self.output), '--font', str(self.font)],
                                capture_output=True, text=True, encoding='utf-8', env=self.environment)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('new', result.stderr)
        self.assertEqual(marker.read_bytes(), before)

    def test_paginate_preserves_ooxml_namespace_prefix_contract(self):
        self.assertTrue(SCRIPT.is_file(), 'qualification entrypoint must exist')
        spec = importlib.util.spec_from_file_location('qualification', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        document = Document()
        for index in range(3):
            document.add_heading(f'图 {index}', level=2)
        path, manifest = self.base / 'namespace.docx', self.base / 'namespace.json'
        document.save(path)
        manifest.write_text(json.dumps({'output': {'sha256': ''}}), encoding='utf-8')
        module._paginate_docx(path, manifest)
        with zipfile.ZipFile(path) as archive:
            root = etree.fromstring(archive.read('word/document.xml'))
        prefixes = root.get('{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable', '').split()
        self.assertTrue(prefixes)
        for prefix in prefixes:
            self.assertIn(prefix, root.nsmap)

    def test_missing_font_fails_without_publishing_passed_report(self):
        self.assertTrue(SCRIPT.is_file(), 'qualification entrypoint must exist')
        output = self.base / 'missing-font-fixture'
        result = subprocess.run([sys.executable, '-X', 'utf8', str(self.entrypoint),
                                 '--output-dir', str(output), '--font', str(self.base / 'missing.ttf')],
                                capture_output=True, text=True, encoding='utf-8', env=self.environment)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('font missing', result.stderr)
        self.assertFalse((output / 'qualification.json').exists())

    def test_timeline_declares_illustrative_milliseconds_and_ticks(self):
        self.assertTrue(SCRIPT.is_file(), 'qualification entrypoint must exist')
        spec = importlib.util.spec_from_file_location('qualification', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        path = self.base / 'timeline-units.svg'
        module.write_diagram(path, *module.DIAGRAMS[1])
        root = ET.parse(path).getroot()
        self.assertTrue(root.find(SVG + 'text').get('font-family').startswith('Noto Sans CJK SC'))
        labels = [node.text for node in root.findall(SVG + 'text')]
        self.assertTrue(any('ms' in label and '示意' in label for label in labels))
        for number in range(0, 51, 10):
            self.assertIn(f'{number} ms', labels)


if __name__ == '__main__':
    unittest.main(verbosity=2)
