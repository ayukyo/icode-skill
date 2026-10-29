#!/usr/bin/env python3
"""Build a .pptx by selecting template slides and editing text or images.

The script preserves all layout, fonts, colors, images and animations of the
source template. It only:
  1. Keeps the slides listed in ``selected_slides`` (in that order),
  2. Replaces the text of the runs identified by ``edits``,
  3. Applies strict ``images`` replacements or empty-region insertions.

Image entries use original template slide numbers, a source path relative to
edits.json, required alt_text, and fit="contain". Specify one of slot_id (from
detail.json pages[].image_slots), address={"shape_id": ...}, or
box_inches=[x, y, width, height]. Replacement targets require exactly one
expected_placeholder (shape name/text) or expected_sha256 (existing image).
Non-picture region shapes must be explicitly declared in image_slots. SVG is
rendered through the sibling DOCX component; --svg-mode dual additionally
preserves its source with a PNG fallback and pending consumer compatibility.
Successful builds write <output-stem>.manifest.json; dry-run preflights inputs.

USAGE
    python3 build_pptx.py <template.pptx> <edits.json> <output.pptx>
                         [--detail detail.json]   # for slot_id resolution
                         [--strict]               # fail if expected_text mismatch
                         [--dry-run]              # show planned edits, do nothing

EDITS JSON SCHEMA
{
  "template_slug": "simple-business",                 # optional, informational
  "selected_slides": [1, 2, 4, 7, 14, 16],            # 1-indexed in template
  "edits": [
    {
      "slide": 1,                                      # 1-indexed in template
      "slot_id": "title_en",                           # resolved via detail.json
      "new_text": "Annual Review"
    },
    {
      "slide": 1,
      "address": {"shape_id": 46, "paragraph": 1, "run": 0},  # explicit
      "expected_text": "工作计划模板",
      "new_text": "2026 年度工作计划"
    }
  ]
}

ADDRESS FORMS
    {"shape_id": <int>, "paragraph": <int>, "run": <int>}   # exact run
    {"shape_id": <int>, "paragraph": <int>}                  # whole paragraph
                                                              # (run 0 keeps fmt,
                                                              # others cleared)
"""
from __future__ import annotations
import argparse
import copy
import hashlib
import json
import math
import re
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


SLIDE_R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


# ------------------------------------------------------------------
# Slot resolution via detail.json
# ------------------------------------------------------------------
def load_slot_index(detail_path: Path | None) -> dict:
    """Return {(slide_number, slot_id): slot_dict}.

    slot_dict includes address + capacity metadata (max_chars, chars_per_line,
    max_lines, wrap, autofit, font_size_pt, level, role) when present so the
    overflow lint can use them.
    """
    if detail_path is None or not detail_path.exists():
        return {}
    detail = json.loads(detail_path.read_text(encoding="utf-8"))
    index: dict = {}
    for page in detail.get("pages", []):
        slide_number = page["slide_number"]
        for slot in page.get("text_slots", []):
            slot.setdefault("expected_text", slot.get("current_text"))
            index[(slide_number, slot["slot_id"])] = slot
    return index


# ------------------------------------------------------------------
# Overflow lint (text-fits-the-box check)
# ------------------------------------------------------------------
def _visual_width(s: str) -> float:
    w = 0.0
    for c in s:
        if "\u4e00" <= c <= "\u9fff" or "\u3000" <= c <= "\u303f" or "\uff00" <= c <= "\uffef":
            w += 1.0
        elif c == " ":
            w += 0.35
        elif c.isascii():
            w += 0.5
        else:
            w += 0.8
    return w


def check_overflow(new_text: str, meta: dict) -> tuple[bool, str]:
    """Return (fits, message). Uses chars_per_line / max_lines from detail.json.

    'fits' is True when the text is expected to stay within the box. Boxes with
    autofit (PowerPoint shrink-to-fit) are always treated as fitting (soft).
    """
    if meta.get("capacity_unknown"):
        return True, ""  # geometry unreliable for this box
    cpl = meta.get("chars_per_line")
    max_lines = meta.get("max_lines")
    cap = meta.get("max_chars")   # already includes the calibration tolerance
    if not cpl or not max_lines or not cap:
        return True, ""  # no capacity data → can't lint
    total_vw = _visual_width(new_text.replace("\n", ""))
    # primary test: visual width vs the (tolerance-adjusted) capacity budget
    if total_vw <= cap:
        return True, ""
    # estimate needed lines for a helpful message
    needed = sum(max(1, math.ceil(_visual_width(seg) / cpl))
                 for seg in (new_text.split("\n") or [new_text]))
    sz = meta.get("font_size_pt")
    msg = (f"视觉宽度 {total_vw:.0f} > 容量 {cap} "
           f"(约需 {needed} 行 / 可容 {max_lines} 行, {sz}pt, 每行≈{cpl}字)")
    if meta.get("autofit"):
        return True, "AUTOFIT " + msg + " — PowerPoint 会自动缩字"
    return False, msg


# ------------------------------------------------------------------
# Shape lookup (recurses into groups)
# ------------------------------------------------------------------
def find_shape(shapes, shape_id: int):
    for shape in shapes:
        if shape.shape_id == shape_id:
            return shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            found = find_shape(shape.shapes, shape_id)
            if found is not None:
                return found
    return None


# ------------------------------------------------------------------
# Text mutation that preserves run-level formatting
# ------------------------------------------------------------------
def apply_edit_to_shape(shape, address: dict, new_text: str, expected: str | None,
                       strict: bool) -> tuple[bool, str]:
    if not shape.has_text_frame:
        return False, f"shape {shape.shape_id} has no text frame"
    tf = shape.text_frame
    pi = address.get("paragraph", 0)
    if pi >= len(tf.paragraphs):
        return False, f"paragraph {pi} out of range (have {len(tf.paragraphs)})"
    para = tf.paragraphs[pi]
    runs = list(para.runs)
    if not runs:
        return False, f"paragraph {pi} has no runs"

    ri = address.get("run")
    if ri is None:
        current = "".join(r.text for r in runs)
        if expected is not None and current != expected and strict:
            return False, (
                f"expected_text mismatch at paragraph level: "
                f"have {current!r}, expected {expected!r}"
            )
        runs[0].text = new_text
        for r in runs[1:]:
            r.text = ""
        return True, f"replaced paragraph (kept run 0 formatting): {current!r} -> {new_text!r}"

    if ri >= len(runs):
        return False, f"run {ri} out of range (have {len(runs)})"
    current = runs[ri].text
    if expected is not None and current != expected and strict:
        return False, (
            f"expected_text mismatch: have {current!r}, expected {expected!r}"
        )
    runs[ri].text = new_text
    return True, f"replaced run: {current!r} -> {new_text!r}"


# ------------------------------------------------------------------
# Slide pruning: keep only selected slides, in order
# ------------------------------------------------------------------
def prune_slides(prs, selected_1indexed: list[int]) -> None:
    sldIdLst = prs.slides._sldIdLst
    sld_ids = list(sldIdLst)
    total = len(sld_ids)
    keep_idx0 = [i - 1 for i in selected_1indexed]
    for i in keep_idx0:
        if i < 0 or i >= total:
            raise ValueError(f"selected slide {i + 1} out of range (1..{total})")

    # Build new order list of element references.
    new_order = [sld_ids[i] for i in keep_idx0]

    # Remove all entries, then re-append in the new order. Keeping the same
    # element instances preserves their r:id and avoids orphaning relationships.
    for sld in list(sldIdLst):
        sldIdLst.remove(sld)
    for sld in new_order:
        sldIdLst.append(sld)


def drop_orphan_slide_parts(path: Path) -> int:
    """Rezip ``path`` dropping slide parts no longer referenced by sldIdLst.

    prune_slides() only removes entries from the slide list; the unselected slide
    parts stay in the package as orphans. Orphans still carry template placeholder
    text (Question N / Vivamus / 项目名称), so they must be stripped from the
    file or the成品 fails the "无占位残留" check even though nothing renders.
    Returns the number of orphan parts removed.
    """
    import zipfile
    from xml.etree import ElementTree

    prs_rels_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    prs_ns = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

    def to_full(target: str) -> str | None:
        """Normalize a rels Target into a package path (or None if not a slide)."""
        if target.startswith("/"):
            target = target.lstrip("/")
        elif not target.startswith("ppt/"):
            target = "ppt/" + target
        return target if re.match(r"ppt/slides/slide\d+\.xml$", target) else None

    with zipfile.ZipFile(path, "r") as z:
        rels = ElementTree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
        target_by_rid = {
            el.get("Id"): el.get("Target")
            for el in rels.findall(prs_rels_ns + "Relationship")
        }
        sld_refs = {
            el.get(r_ns + "id")
            for el in ElementTree.fromstring(
                z.read("ppt/presentation.xml")).iter(prs_ns + "sldId")
        }
        referenced = {
            p for rid in sld_refs
            if (p := to_full(target_by_rid[rid])) is not None
        }
        drop = {
            n for n in z.namelist()
            if re.match(r"ppt/slides/slide\d+\.xml$", n) and n not in referenced
        } | {
            n for n in z.namelist()
            if re.match(r"ppt/slides/_rels/slide\d+\.xml\.rels$", n)
               and n.replace("_rels/", "").replace(".rels", "") not in referenced
        }
        if not drop:
            return 0
        kept = [n for n in z.namelist() if n not in drop]

    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n in kept:
            z.writestr(n, zipfile.ZipFile(path, "r").read(n))
    tmp.replace(path)
    return len(drop)


# ------------------------------------------------------------------
# Main driver
# ------------------------------------------------------------------
def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _keys(value, allowed: set[str], label: str) -> None:
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be an object")
    unknown = set(value) - allowed
    if unknown:
        raise ValueError(f"unknown {label} fields: {sorted(unknown)}")


def _positive_int(value, label: str) -> int:
    if type(value) is not int or value < 1:
        raise ValueError(f"{label} must be a positive integer")
    return value


def _docx_component(name: str):
    # The PPT tool reuses the sibling DOCX SVG components; ordinary PNG/JPEG
    # edits never import them or require an SVG rendering installation.
    import importlib
    sibling = Path(__file__).resolve().parents[2] / "docx"
    if str(sibling) not in sys.path:
        sys.path.insert(0, str(sibling))
    try:
        return importlib.import_module(name)
    except ImportError as exc:
        raise ValueError(f"SVG requires tools/docx/{name}.py and its dependencies: {exc}") from exc


def _intersects(a, b) -> bool:
    return (a[0] < b[0] + b[2] and b[0] < a[0] + a[2]
            and a[1] < b[1] + b[3] and b[1] < a[1] + a[3])


def _plan_images(prs, images, selected, detail_path, args) -> list[dict]:
    if not isinstance(images, list):
        raise ValueError("images must be a list")
    slots = {}
    if detail_path is not None:
        detail = json.loads(detail_path.read_text(encoding="utf-8"))
        for page in detail.get("pages", []):
            for slot in page.get("image_slots", []):
                key = (page["slide_number"], slot["slot_id"])
                if key in slots:
                    raise ValueError(f"duplicate image slot {key}")
                slots[key] = slot
    plans, targets, insert_boxes = [], set(), {}
    for entry in images:
        _keys(entry, {"slide", "source", "alt_text", "fit", "slot_id", "address",
                      "box_inches", "expected_placeholder", "expected_sha256"}, "image")
        slide_num = _positive_int(entry.get("slide"), "image slide")
        if slide_num not in selected:
            raise ValueError(f"image slide {slide_num} is not selected")
        if entry.get("fit", "contain") != "contain":
            raise ValueError("image fit must be contain")
        if not isinstance(entry.get("alt_text"), str) or not entry["alt_text"].strip():
            raise ValueError("image alt_text must be a nonempty string")
        # lxml rejects control characters only at assignment time; validate here
        # so dry-run offers the same input guarantee as a full build.
        from lxml import etree
        etree.Element("alt").set("descr", entry["alt_text"])
        if sum(key in entry for key in ("slot_id", "address", "box_inches")) != 1:
            raise ValueError("image requires exactly one of slot_id, address, box_inches")
        if not isinstance(entry.get("source"), str) or not entry["source"].strip():
            raise ValueError("image source must be a nonempty path")
        source = (args.edits.parent / entry["source"]).resolve()
        if not source.is_file():
            raise ValueError(f"image source is not a file: {source}")
        slide = prs.slides[slide_num - 1]
        shape, meta = None, {}
        if "slot_id" in entry:
            if not isinstance(entry["slot_id"], str):
                raise ValueError("image slot_id must be a string")
            meta = slots.get((slide_num, entry["slot_id"]))
            if meta is None:
                raise ValueError(f"image slot_id {entry['slot_id']!r} not found on slide {slide_num}")
            address = meta.get("address")
        else:
            address = entry.get("address")
            if isinstance(address, dict):
                matches = [value for (number, _), value in slots.items()
                           if number == slide_num and value.get("address") == address]
                if len(matches) > 1:
                    raise ValueError("image address has ambiguous detail.json slot declarations")
                if matches:
                    meta = matches[0]
        if address is not None:
            _keys(address, {"shape_id"}, "image address")
            shape_id = _positive_int(address.get("shape_id"), "image shape_id")
            shape = find_shape(slide.shapes, shape_id)
            if shape is None:
                raise ValueError(f"image shape_id {shape_id} not found on slide {slide_num}")
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                raise ValueError("image region cannot replace a group and discard its children")
            if (slide_num, shape_id) in targets:
                raise ValueError(f"duplicate image target slide {slide_num} shape {shape_id}")
            targets.add((slide_num, shape_id))
            is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            if not is_picture and not meta:
                raise ValueError("non-picture replacement requires a declared detail.json image slot")
            placeholder = entry.get("expected_placeholder", meta.get("expected_placeholder"))
            expected_hash = entry.get("expected_sha256", meta.get("expected_sha256"))
            if (placeholder is None) == (expected_hash is None):
                raise ValueError("image replacement requires exactly one expected_placeholder or expected_sha256")
            if placeholder is not None:
                if not isinstance(placeholder, str) or not placeholder:
                    raise ValueError("expected_placeholder must be a nonempty string")
                actual = {shape.name}
                if shape.has_text_frame:
                    actual.add(shape.text)
                if placeholder not in actual:
                    raise ValueError(f"expected_placeholder mismatch for shape {shape_id}")
            if expected_hash is not None:
                if not isinstance(expected_hash, str) or not re.fullmatch(r"[0-9a-fA-F]{64}", expected_hash):
                    raise ValueError("expected_sha256 must contain 64 hexadecimal characters")
                if not is_picture or hashlib.sha256(shape.image.blob).hexdigest() != expected_hash.lower():
                    raise ValueError(f"expected_sha256 mismatch for shape {shape_id}")
            box = (int(shape.left), int(shape.top), int(shape.width), int(shape.height))
            if box[2] <= 0 or box[3] <= 0:
                raise ValueError("image slot geometry must have positive width and height")
        else:
            if "expected_placeholder" in entry or "expected_sha256" in entry:
                raise ValueError("free box insertion cannot have replacement expectations")
            raw_box = entry.get("box_inches")
            if (not isinstance(raw_box, list) or len(raw_box) != 4 or
                    any(type(x) not in (int, float) or not math.isfinite(x) for x in raw_box)):
                raise ValueError("box_inches must contain four finite numbers")
            x, y, w, h = raw_box
            box = tuple(int(Inches(v)) for v in raw_box)
            if (x < 0 or y < 0 or w <= 0 or h <= 0 or box[2] <= 0 or box[3] <= 0
                    or box[0] + box[2] > prs.slide_width or box[1] + box[3] > prs.slide_height):
                raise ValueError("box_inches is outside slide bounds or has nonpositive size")
            occupied = [(s.left, s.top, s.width, s.height) for s in slide.shapes]
            occupied += insert_boxes.setdefault(slide_num, [])
            if any(_intersects(box, other) for other in occupied):
                raise ValueError("box_inches overlaps an existing shape or planned image")
            insert_boxes[slide_num].append(box)
        suffix = source.suffix.lower()
        diagram = None
        if suffix == ".svg":
            renderer = _docx_component("diagram_svg")
            if getattr(args, "svg_mode", "raster") == "dual":
                _docx_component("diagram_ooxml")
            # Published provenance must survive system temporary-cache cleanup.
            # Dry-run has no delivered package and keeps its rendering temporary.
            cache_dir = (Path(tempfile.mkdtemp(prefix="icode-ppt-svg-")) if args.dry_run else
                         args.output.parent / (args.output.stem + ".assets") / "svg")
            diagram = renderer.prepare_svg(source, cache_dir,
                       width_inches=box[2] / Inches(1), dpi=getattr(args, "svg_dpi", 192),
                       font_files=getattr(args, "svg_font", None))
            raster = Path(diagram["fallback_png"])
        elif suffix in (".png", ".jpg", ".jpeg"):
            raster = source
        else:
            raise ValueError(f"unsupported image format: {suffix}")
        from PIL import Image
        with Image.open(raster) as image:
            if image.format not in ("PNG", "JPEG"):
                raise ValueError("image content must be PNG or JPEG")
            iw, ih = image.size
            image.verify()
        scale = min(box[2] / iw, box[3] / ih)
        width, height = max(1, int(iw * scale)), max(1, int(ih * scale))
        geometry = {"left": box[0] + (box[2] - width) // 2,
                    "top": box[1] + (box[3] - height) // 2, "width": width, "height": height}
        record = {"original_slide": slide_num, "output_slide": selected.index(slide_num) + 1,
                  "source": str(source), "sha256": _sha256(source), "source_format": suffix[1:],
                  "embedded_sha256": _sha256(raster), "alt_text": entry["alt_text"],
                  "geometry": geometry, "slot_box": list(box), "fit": "contain",
                  "slot_id": entry.get("slot_id"), "shape_id": shape.shape_id if shape else None,
                  "embedded_format": "png_fallback" if diagram else suffix[1:],
                  "vector_preserved": False}
        if diagram:
            record.update(diagram)
            record.update(geometry=geometry, original_slide=slide_num,
                          output_slide=selected.index(slide_num) + 1, alt_text=entry["alt_text"],
                          shape_id=shape.shape_id if shape else None)
            record["source"] = str(source)
            record["fallback_png"] = str(raster.resolve())
            if getattr(args, "svg_mode", "raster") == "dual":
                record.update(embedded_format="svg_with_png_fallback", vector_preserved=True,
                              consumer_compatibility="pending", diagram_status="native_vector_pending")
        plans.append({"slide": slide, "shape": shape, "raster": raster, "record": record,
                      "is_diagram": diagram is not None})
    if getattr(args, "svg_mode", "raster") == "dual":
        by_fallback = {}
        for plan in plans:
            if not plan["is_diagram"]:
                continue
            record = plan["record"]
            key = record["fallback_sha256"]
            if key in by_fallback and by_fallback[key] != record["sha256"]:
                raise ValueError("ambiguous SVG sources share one PNG fallback")
            by_fallback[key] = record["sha256"]
    protected = {args.template.resolve(), args.edits.resolve()}
    if detail_path is not None:
        protected.add(detail_path.resolve())
    protected.update(Path(p["record"]["source"]).resolve() for p in plans)
    protected.update(Path(font).resolve() for font in (getattr(args, "svg_font", None) or []))
    outputs = {args.output.resolve(), args.output.with_suffix(".manifest.json").resolve()}
    if outputs & protected:
        raise ValueError("output or manifest would overwrite an input")
    return plans


def _apply_images(plans) -> None:
    for plan in plans:
        record, old = plan["record"], plan["shape"]
        geo = record["geometry"]
        picture = plan["slide"].shapes.add_picture(str(plan["raster"]), **geo)
        if old is not None:
            # Keep the old identity and exact sibling position, including when
            # the target lives in a group. Slide relationships still own media.
            props = picture._element.nvPicPr.cNvPr
            old_props = old._element.xpath(".//p:cNvPr")[0]
            props.attrib.update(old_props.attrib)
            new_element = picture._element
            new_element.getparent().remove(new_element)
            old._element.getparent().replace(old._element, new_element)
        picture._element.nvPicPr.cNvPr.set("descr", record["alt_text"])
        record["shape_id"] = picture.shape_id


def run(args) -> int:
    edits_spec = json.loads(args.edits.read_text(encoding="utf-8"))
    _keys(edits_spec, {"template_slug", "selected_slides", "edits", "images"}, "edits specification")
    selected = edits_spec.get("selected_slides", [])
    raw_edits = edits_spec.get("edits", [])
    if not isinstance(selected, list) or not isinstance(raw_edits, list):
        raise ValueError("selected_slides and edits must be lists")
    slot_index = load_slot_index(args.detail)

    prs = Presentation(args.template)
    n_slides = len(prs.slides)
    if not selected:
        selected = list(range(1, n_slides + 1))
    for s in selected:
        _positive_int(s, "selected slide")
        if s < 1 or s > n_slides:
            print(f"ERROR: selected slide {s} out of range 1..{n_slides}", file=sys.stderr)
            return 2
    if len(set(selected)) != len(selected):
        raise ValueError("selected_slides cannot contain duplicates")
    image_plans = _plan_images(prs, edits_spec.get("images", []), selected, args.detail, args)

    # Resolve and group edits by slide_number (in TEMPLATE order, pre-prune)
    planned: list[dict] = []
    for e in raw_edits:
        _keys(e, {"slide", "slot_id", "address", "expected_text", "new_text"}, "text edit")
        slide_num = e["slide"]
        _positive_int(slide_num, "text edit slide")
        if slide_num not in selected:
            raise ValueError(f"text edit slide {slide_num} is not selected")
        if not isinstance(e.get("new_text"), str):
            raise ValueError("new_text must be a string")
        meta = None
        if "address" in e:
            addr = dict(e["address"])
            expected = e.get("expected_text") or addr.get("expected_text")
        else:
            slot_id = e.get("slot_id")
            if slot_id is None:
                print(f"ERROR: edit {e} missing both address and slot_id", file=sys.stderr)
                return 2
            key = (slide_num, slot_id)
            if key not in slot_index:
                print(f"ERROR: slot_id {slot_id!r} not found for slide {slide_num} in detail.json",
                      file=sys.stderr)
                return 2
            meta = slot_index[key]
            addr = dict(meta["address"])
            expected = e.get("expected_text") or meta.get("expected_text")
        planned.append({
            "slide": slide_num,
            "address": addr,
            "expected": expected,
            "new_text": e["new_text"],
            "meta": meta,
            "raw": e,
        })
        _keys(addr, {"shape_id", "paragraph", "run", "expected_text"}, "text address")
        _positive_int(addr.get("shape_id"), "text shape_id")
        for index_key in ("paragraph", "run"):
            if index_key in addr and (type(addr[index_key]) is not int or addr[index_key] < 0):
                raise ValueError(f"text {index_key} must be a nonnegative integer")
        if any(p["record"]["original_slide"] == slide_num and
               p["record"]["shape_id"] == addr["shape_id"] for p in image_plans):
            raise ValueError("text and image edits cannot target the same shape")

    # Apply edits BEFORE pruning so slide indices match the template
    overflow_issues: list[str] = []
    ellipsis_issues: list[str] = []
    for p in planned:
        slide = prs.slides[p["slide"] - 1]
        shape_id = p["address"].get("shape_id")
        shape = find_shape(slide.shapes, shape_id)
        if shape is None:
            msg = f"slide {p['slide']}: shape_id {shape_id} not found"
            if args.strict:
                print(f"ERROR: {msg}", file=sys.stderr)
                return 3
            print(f"WARN:  {msg}", file=sys.stderr)
            continue
        ok, info = apply_edit_to_shape(
            shape, p["address"], p["new_text"], p["expected"], args.strict
        )
        prefix = "OK  " if ok else ("ERR " if args.strict else "WARN")
        print(f"{prefix} slide {p['slide']:>3} shape {shape_id}: {info}")
        if not ok and args.strict:
            return 4

        # Overflow check — advisory only. max_chars is a soft hint, not a hard
        # limit; we never block or truncate. We only print a gentle note so the
        # author can decide whether to rephrase more concisely.
        if not args.no_lint and p["meta"]:
            fits, omsg = check_overflow(p["new_text"], p["meta"])
            if omsg and not fits:
                slot_id = p["meta"].get("slot_id", "?")
                line = (f"slide {p['slide']:>3} {slot_id}: 可能偏长 {omsg} "
                        f"-> {p['new_text'][:24]!r}")
                overflow_issues.append(line)
                print(f"note {line}", file=sys.stderr)
            elif omsg and fits:  # autofit soft note
                print(f"note  slide {p['slide']:>3}: {omsg}", file=sys.stderr)

        # Catch accidental ellipsis-truncation, which looks worse than a slight
        # overflow. Advisory only.
        stripped = p["new_text"].rstrip()
        if stripped.endswith(("...", "..", "…", "……", "等等", "等。")):
            ellipsis_issues.append(
                f"slide {p['slide']:>3} {(p['meta'] or {}).get('slot_id', '?')}: "
                f"结尾疑似省略号截断 -> {p['new_text'][-16:]!r}")

    if overflow_issues:
        print(f"\n提示：{len(overflow_issues)} 处文字偏长（仅提示，不影响保存）。"
              f"可酌情用更精炼的措辞重写；但**不要截断加省略号**，也不要改字号——"
              f"宁可轻微超出。", file=sys.stderr)

    if ellipsis_issues:
        print(f"\n⚠️ {len(ellipsis_issues)} 处文字结尾是省略号，很可能是为凑长度而截断"
              f"（这是最差的结果）。请改成完整、精炼的表达：", file=sys.stderr)
        for line in ellipsis_issues:
            print(f"  - {line}", file=sys.stderr)

    if args.dry_run:
        print(f"Preflight passed: {len(planned)} text edits, {len(image_plans)} images")
        return 0
    _apply_images(image_plans)
    prune_slides(prs, selected)
    diagrams = [p["record"] for p in image_plans if p["is_diagram"]]
    manifest = {"schema_version": 1, "template": {"path": str(args.template.resolve()), "sha256": _sha256(args.template)},
                "input": {"path": str(args.edits.resolve()), "sha256": _sha256(args.edits)},
                "output": {"path": str(args.output.resolve())},
                "selected_slides": selected, "images": [p["record"] for p in image_plans],
                "diagrams": diagrams, "visual_review": "pending"}
    # Publish only after package generation and optional dual embedding succeed.
    # In particular an SVG/helper failure must leave an existing output intact.
    with tempfile.TemporaryDirectory(prefix=".ppt-build-", dir=args.output.parent) as staging:
        package = Path(staging) / args.output.name
        prs.save(package)
        orphaned = drop_orphan_slide_parts(package)
        if diagrams and getattr(args, "svg_mode", "raster") == "dual":
            helper = _docx_component("diagram_ooxml")
            helper.embed_svg_sources(package, diagrams)
            checks = helper.inspect_diagram_package(package, diagrams)
            if not all(checks.values()):
                raise ValueError(f"SVG package checks failed: {checks}")
        manifest["output"]["sha256"] = _sha256(package)
        manifest_tmp = Path(staging) / "manifest.json"
        manifest_tmp.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        package.replace(args.output)
        manifest_tmp.replace(args.output.with_suffix(".manifest.json"))
    note = f"\nSaved {args.output} with {len(selected)} slides and {len(planned)} edits"
    if orphaned:
        note += f" (清理 {orphaned} 个未引用孤儿 slide 部件)"
    print(note)
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("template", type=Path)
    ap.add_argument("edits", type=Path)
    ap.add_argument("output", type=Path)
    ap.add_argument("--detail", type=Path, default=None)
    ap.add_argument("--no-lint", action="store_true",
                    help="disable the text-overflow check")
    ap.add_argument("--strict", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--svg-font", type=Path, action="append", default=None)
    ap.add_argument("--svg-dpi", type=int, default=192)
    ap.add_argument("--svg-mode", choices=("raster", "dual"), default="raster")
    args = ap.parse_args()

    if args.detail is None:
        # Try sibling detail.json next to the template
        guess = args.template.with_name("detail.json")
        if guess.exists():
            args.detail = guess
            print(f"(auto-detected detail.json: {guess})", file=sys.stderr)

    try:
        return run(args)
    except (ValueError, OSError, KeyError, TypeError, IndexError, RuntimeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
