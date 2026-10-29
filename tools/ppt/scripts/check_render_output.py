#!/usr/bin/env python3
"""Deterministically validate PPT slide previews without visual-model input."""
from __future__ import annotations

import argparse
import hashlib
import importlib
import json
import re
import struct
import sys
import zipfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.exc import PackageNotFoundError


PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
PREVIEW_NAME = re.compile(r"^slide-(\d+)\.png$")


def _png_info(path: Path) -> dict[str, Any]:
    item: dict[str, Any] = {"path": str(path), "size": 0, "width": None, "height": None}
    try:
        if path.is_symlink() or not path.is_file():
            raise ValueError("not_regular_file")
        item["size"] = path.stat().st_size
        header = path.read_bytes()[:24]
        if len(header) < 24 or header[:8] != PNG_SIGNATURE or header[12:16] != b"IHDR":
            raise ValueError("invalid_png_header")
        width, height = struct.unpack(">II", header[16:24])
        if width <= 0 or height <= 0:
            raise ValueError("invalid_dimensions")
        item["width"] = width
        item["height"] = height
        item["valid"] = True
    except (OSError, ValueError, struct.error) as exc:
        item["valid"] = False
        item["error"] = str(exc)
    return item


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _find_shape(shapes, shape_id):
    for shape in shapes:
        if shape.shape_id == shape_id:
            return shape
        if hasattr(shape, "shapes"):
            found = _find_shape(shape.shapes, shape_id)
            if found is not None:
                return found
    return None


def _check_manifest(pptx_path: Path, manifest_path: Path) -> tuple[dict, list[str]]:
    checks, errors = {}, []

    def check(name, passed):
        checks[name] = bool(passed)
        if not passed:
            errors.append(f"manifest {name} mismatch")

    try:
        manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
        check("output_hash", _sha256(pptx_path) == manifest["output"]["sha256"])
        for name in ("template", "input"):
            check(f"{name}_hash", _sha256(Path(manifest[name]["path"])) == manifest[name]["sha256"])
        prs = Presentation(str(pptx_path))
        images = manifest["images"]
        diagrams = manifest["diagrams"]
        if not isinstance(images, list) or not isinstance(diagrams, list):
            raise ValueError("images and diagrams must be lists")
        check("diagram_records", diagrams == [r for r in images if r["source_format"] == "svg"])
        check("selected_slide_count", len(manifest["selected_slides"]) == len(prs.slides))
        for i, record in enumerate(images):
            prefix = f"image_{i}"
            source = Path(record["source"])
            check(f"{prefix}_source_hash", _sha256(source) == record["sha256"])
            if record["source_format"] == "svg":
                check(f"{prefix}_fallback_hash", _sha256(Path(record["fallback_png"])) == record["fallback_sha256"])
                check(f"{prefix}_fallback_embedded", record["fallback_sha256"] == record["embedded_sha256"])
            slide_number = record["output_slide"]
            if type(slide_number) is not int or not 1 <= slide_number <= len(prs.slides):
                raise ValueError(f"{prefix} output_slide is out of range")
            check(f"{prefix}_slide_mapping", manifest["selected_slides"][slide_number - 1] == record["original_slide"])
            shape = _find_shape(prs.slides[slide_number - 1].shapes, record["shape_id"])
            if shape is None or not hasattr(shape, "image"):
                raise ValueError(f"{prefix} picture shape is missing")
            check(f"{prefix}_embedded_hash", hashlib.sha256(shape.image.blob).hexdigest() == record["embedded_sha256"])
            geometry = {key: int(getattr(shape, key)) for key in ("left", "top", "width", "height")}
            check(f"{prefix}_geometry", geometry == record["geometry"])
            check(f"{prefix}_alt_text", shape._element.nvPicPr.cNvPr.get("descr") == record["alt_text"])
        if any(r.get("embedded_format") == "svg_with_png_fallback" for r in diagrams):
            sibling = Path(__file__).resolve().parents[2] / "docx"
            if str(sibling) not in sys.path:
                sys.path.insert(0, str(sibling))
            try:
                helper = importlib.import_module("diagram_ooxml")
            except ImportError as exc:
                raise ValueError(f"SVG verification requires tools/docx/diagram_ooxml.py: {exc}") from exc
            for name, passed in helper.inspect_diagram_package(pptx_path, diagrams).items():
                check(f"diagram_package_{name}", passed)
    except (OSError, ValueError, KeyError, TypeError, IndexError, AttributeError,
            PackageNotFoundError, zipfile.BadZipFile) as exc:
        errors.append(f"manifest validation failed: {exc}")
    return checks, errors


def check_render_output(pptx_path: Path, preview_dir: Path, manifest_path: Path | None = None) -> dict[str, Any]:
    errors: list[str] = []
    manifest_checks = None
    if manifest_path is not None:
        manifest_checks, manifest_errors = _check_manifest(pptx_path, manifest_path)
        errors.extend(manifest_errors)
    try:
        if pptx_path.is_symlink() or not pptx_path.is_file():
            raise ValueError(f"PPTX 不是普通文件: {pptx_path}")
        slide_count = len(Presentation(str(pptx_path)).slides)
    except (OSError, ValueError, KeyError, PackageNotFoundError, zipfile.BadZipFile) as exc:
        return {
            "status": "failed",
            "pptx": str(pptx_path),
            "preview_dir": str(preview_dir),
            "errors": errors + [str(exc)],
            **({"manifest_checks": manifest_checks} if manifest_path is not None else {}),
        }

    if not preview_dir.is_dir() or preview_dir.is_symlink():
        return {
            "status": "failed",
            "pptx": str(pptx_path),
            "preview_dir": str(preview_dir),
            "slide_count": slide_count,
            "preview_count": 0,
            "missing": [f"slide-{number}.png" for number in range(1, slide_count + 1)],
            "extras": [],
            "invalid": [],
            "previews": [],
            "errors": errors + ["预览目录不存在或不是普通目录"],
            **({"manifest_checks": manifest_checks} if manifest_path is not None else {}),
        }

    actual_paths = sorted(preview_dir.glob("slide-*.png"), key=lambda path: path.name)
    previews: list[dict[str, Any]] = []
    page_files: dict[int, list[str]] = {}
    invalid_names: list[str] = []
    for path in actual_paths:
        match = PREVIEW_NAME.fullmatch(path.name)
        item = _png_info(path)
        if match:
            page = int(match.group(1))
            item["page"] = page
            page_files.setdefault(page, []).append(path.name)
        else:
            item["page"] = None
            item["valid"] = False
            item["error"] = "invalid_preview_name"
            invalid_names.append(path.name)
        previews.append(item)
    previews.sort(key=lambda item: (item["page"] is None, item["page"] or 0, item["path"]))
    missing = [f"slide-{number}.png" for number in range(1, slide_count + 1) if number not in page_files]
    extras = sorted(
        name for page, names in page_files.items() if page < 1 or page > slide_count for name in names
    )
    duplicate_pages = [
        {"page": page, "files": sorted(names)}
        for page, names in sorted(page_files.items())
        if len(names) > 1
    ]
    invalid = sorted(
        set(invalid_names)
        | {Path(item["path"]).name for item in previews if not item.get("valid")}
    )
    if slide_count == 0:
        errors.append("PPTX 不含幻灯片")
    if missing:
        errors.append("缺少渲染页")
    if extras:
        errors.append("存在不属于当前 PPTX 的额外预览页")
    if invalid:
        errors.append("预览 PNG 签名或尺寸无效")
    if duplicate_pages:
        errors.append("同一页存在重复预览文件")
    return {
        "status": "passed" if not errors else "failed",
        "pptx": str(pptx_path),
        "preview_dir": str(preview_dir),
        "slide_count": slide_count,
        "preview_count": len(actual_paths),
        "missing": missing,
        "extras": extras,
        "duplicate_pages": duplicate_pages,
        "invalid": invalid,
        "previews": previews,
        "errors": errors,
        **({"manifest_checks": manifest_checks} if manifest_path is not None else {}),
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="检查 PPTX 页数与 slide-N.png 数量、PNG 签名和尺寸；不做视觉语义判断。"
    )
    parser.add_argument("pptx", type=Path)
    parser.add_argument("preview_dir", type=Path)
    parser.add_argument("--manifest", type=Path)
    args = parser.parse_args()
    result = check_render_output(args.pptx, args.preview_dir, args.manifest)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    sys.exit(main())
